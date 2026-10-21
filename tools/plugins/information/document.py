"""Document reading: PDF, DOCX, XLSX, CSV, code summarization.

Uses free, offline-capable libraries:
- PDF: pypdf (BSD)
- PPTX: python-pptx
- DOCX: python-docx
- XLSX: openpyxl
- CSV: built-in csv module
- Code: plain text files
"""
from __future__ import annotations

import asyncio
import csv
import logging
from pathlib import Path

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import pptx
except ImportError:
    pptx = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".pptx": "pptx",
    ".csv": "csv",
    ".py": "code",
    ".js": "code",
    ".ts": "code",
    ".sh": "code",
    ".md": "code",
    ".txt": "code",
    ".json": "code",
    ".toml": "code",
    ".yaml": "code",
    ".yml": "code",
}


async def read_document(filepath: str | Path) -> dict:
    path = Path(filepath)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")
    result = {"type": SUPPORTED_EXTENSIONS[ext], "filepath": str(path)}
    if ext == ".pdf" and pypdf:
        result.update(await _read_pdf(path))
    elif ext == ".docx" and DocxDocument:
        result.update(await _read_docx(path))
    elif ext == ".pptx" and pptx:
        result.update(await _read_pptx(path))
    elif ext == ".xlsx" and openpyxl:
        result.update(await _read_xlsx(path))
    elif ext == ".csv":
        result.update(await _read_csv(path))
    elif SUPPORTED_EXTENSIONS[ext] == "code":
        result.update(await _read_code(path))
    else:
        raise ValueError(f"Library not installed for {ext}")
    return result


async def _read_pdf(path: Path) -> dict:
    def _sync():
        reader = pypdf.PdfReader(path)
        pages = []
        for i, page in enumerate(reader.pages):
            pages.append({"page": i + 1, "text": page.extract_text() or ""})
        text = "\n".join(item["text"] for item in pages)
        return {"pages": pages, "text": text, "page_count": len(pages)}
    return await asyncio.to_thread(_sync)


async def _read_docx(path: Path) -> dict:
    def _sync():
        doc = DocxDocument(path)
        from docx.table import Table
        text_parts = []
        tables = []
        for block in doc.iter_inner_content():
            if isinstance(block, Table):
                table_data = [[cell.text for cell in row.cells] for row in block.rows]
                tables.append(table_data)
                text_parts.extend(" | ".join(row) for row in table_data)
            else:
                text_parts.append(block.text)
        return {"text": "\n".join(text_parts), "tables": tables}
    return await asyncio.to_thread(_sync)


async def _read_pptx(path: Path) -> dict:
    def _sync():
        deck = pptx.Presentation(path)
        slides = []
        for i, slide in enumerate(deck.slides):
            parts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text]
            slides.append({"slide": i + 1, "text": "\n".join(parts)})
        return {"slides": slides, "text": "\n".join(item["text"] for item in slides), "page_count": len(slides)}
    return await asyncio.to_thread(_sync)


async def _read_xlsx(path: Path) -> dict:
    def _sync():
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheets = {}
        texts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_text = ""
            sheet_rows = []
            for row in ws.iter_rows(min_row=1, values_only=True):
                sheet_rows.append([str(c) if c is not None else "" for c in row])
                sheet_text += ", ".join(str(c) if c is not None else "" for c in row) + "\n"
            sheets[sheet_name] = {"rows": sheet_rows, "text": sheet_text}
            texts.append(sheet_text)
        return {"sheets": sheets, "text": "\n".join(texts)}
    return await asyncio.to_thread(_sync)


async def _read_csv(path: Path) -> dict:
    def _sync():
        rows = []
        text = ""
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            if header:
                text += ", ".join(header) + "\n"
            for row in reader:
                rows.append(row)
                text += ", ".join(row) + "\n"
        return {"rows": rows, "text": text, "has_header": header is not None}
    return await asyncio.to_thread(_sync)


async def _read_code(path: Path) -> dict:
    def _sync():
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = path.read_text(encoding="latin-1")
        lines = text.splitlines()
        return {"text": text, "line_count": len(lines), "char_count": len(text)}
    return await asyncio.to_thread(_sync)


async def summarize_text(text: str, max_lines: int = 100) -> str:
    if not text:
        return ""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return ""
    if len(lines) <= max_lines:
        return "\n".join(lines)
    keywords = {"error", "warning", "failure", "success", "result", "summary",
                "total", "average", "max", "min", "top", "best"}
    important = [line for line in lines if any(k in line.lower() for k in keywords)]
    summary = lines[:3]
    if important:
        summary.extend(important[:3])
    summary.extend(lines[-3:])
    seen: set[str] = set()
    deduped = []
    for line in summary:
        if line not in seen:
            seen.add(line)
            deduped.append(line)
    return "\n".join(deduped)


async def can_read(filepath: str | Path) -> bool:
    path = Path(filepath)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        return False
    if ext == ".pdf" and pypdf is None:
        return False
    if ext == ".docx" and DocxDocument is None:
        return False
    if ext == ".pptx" and pptx is None:
        return False
    if ext == ".xlsx" and openpyxl is None:
        return False
    return True
