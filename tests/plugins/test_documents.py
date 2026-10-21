"""Document reading (pypdf, python-docx, openpyxl, python-pptx) and the folder-limited read_document tool (PLAN M2-3)."""
import asyncio
import os
import time
from pathlib import Path

import pytest

from tools.plugins.information.document import read_document
from tools.plugins.system import actions
from tools.plugins.toolkit import TOOLS, ActionQueue, PluginContext, Toolbox


def _pdf(path: Path, text: str) -> None:
    """A minimal one-page PDF with one line of Helvetica text."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))


@pytest.fixture()
def home(tmp_path, monkeypatch):
    for name in ("Desktop", "Documents", "Downloads", "Pictures", "Music", "Videos"):
        (tmp_path / name).mkdir()
    monkeypatch.setattr(actions, "HOME", tmp_path)
    monkeypatch.setattr(actions, "DOCUMENT_ROOTS", [tmp_path / n for n in ("Desktop", "Documents", "Downloads")])
    monkeypatch.setattr(actions, "SEARCH_ROOTS", [tmp_path / n for n in ("Desktop", "Documents", "Downloads",
                                                                          "Pictures", "Music", "Videos")])
    return tmp_path


def test_each_format_is_extracted(home):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    folder = home / "Documents"
    _pdf(folder / "invoice.pdf", "Invoice total INR 4,250.00")
    doc = Document()
    doc.add_paragraph("Meeting notes for Tuesday")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Invoice total"
    table.cell(0, 1).text = "INR 4,250.00"
    doc.add_paragraph("End of invoice")
    doc.save(folder / "notes.docx")
    book = Workbook()
    book.active.append(["Item", "Amount"])
    book.active.append(["Rent", 18000])
    book.save(folder / "budget.xlsx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Quarterly plan"
    deck.save(folder / "plan.pptx")
    (folder / "todo.txt").write_text("buy milk", encoding="utf-8")

    def read(name):
        return asyncio.run(read_document(folder / name))

    pdf = read("invoice.pdf")
    assert "Invoice total INR 4,250.00" in pdf["text"] and pdf["page_count"] == 1
    assert "Meeting notes for Tuesday" in read("notes.docx")["text"]
    doc_text = read("notes.docx")["text"]
    assert "Invoice total | INR 4,250.00" in doc_text
    assert doc_text.index("Tuesday") < doc_text.index("Invoice total") < doc_text.index("End of invoice")
    assert "Rent, 18000" in read("budget.xlsx")["text"]
    deck_data = read("plan.pptx")
    assert "Quarterly plan" in deck_data["text"] and deck_data["page_count"] == 1
    assert read("todo.txt")["text"] == "buy milk"
    with pytest.raises(ValueError):
        asyncio.run(read_document(folder / "old.doc"))  # not readable by python-docx: refused, not guessed


def _toolbox(home, safety="ask_risky"):
    class Profile:
        def get(self, _user):
            return {"safety": safety}

    context = PluginContext("web", None, None, "owner", actions=ActionQueue(home).view(), profile=Profile(), is_admin=True)
    return Toolbox(TOOLS, context)


def test_read_document_tool_reads_wraps_and_truncates(home, monkeypatch):
    monkeypatch.setenv("DISPLAY", ":test")
    _pdf(home / "Downloads" / "invoice.pdf", "Invoice total INR 4,250.00")
    (home / "Documents" / "long.txt").write_text("word " * 5000, encoding="utf-8")
    box = _toolbox(home)
    reply = asyncio.run(box.call("read_document", {"path": str(Path("Downloads") / "invoice.pdf")}))
    assert reply.startswith("invoice.pdf, 1 pages:") and "INR 4,250.00" in reply
    assert "<<UNTRUSTED DOCUMENT CONTENT" in reply and reply.rstrip().endswith("<<END>>")
    long = asyncio.run(box.call("read_document", {"path": str(home / "Documents" / "long.txt")}))
    assert "first part only" in long and "[truncated]" in long and len(long) <= 4000


def test_read_document_refuses_outside_paths(home, monkeypatch, tmp_path_factory):
    monkeypatch.setenv("DISPLAY", ":test")
    outside = tmp_path_factory.mktemp("elsewhere") / "secret.txt"
    outside.write_text("api key", encoding="utf-8")
    (home / "Pictures" / "note.txt").write_text("not a document folder", encoding="utf-8")
    box = _toolbox(home)
    for path in (str(outside), "Pictures/note.txt", "Documents/../../" + outside.parent.name + "/secret.txt",
                 "Documents/missing.pdf"):
        reply = asyncio.run(box.call("read_document", {"path": path}))
        assert reply.startswith("Invalid input for read_document"), (path, reply)
        assert "api key" not in reply


def test_read_document_asks_first_by_default(home, monkeypatch):
    monkeypatch.setenv("DISPLAY", ":test")
    (home / "Documents" / "a.txt").write_text("x", encoding="utf-8")
    box = _toolbox(home, safety="ask_all")
    assert "Approval requested" in asyncio.run(box.call("read_document", {"path": "Documents/a.txt"}))


def test_search_files_ext_filter_newest_first(home):
    older, newer = home / "Downloads" / "invoice-march.pdf", home / "Documents" / "invoice-april.pdf"
    for path in (older, newer, home / "Downloads" / "invoice-april.docx"):
        path.write_bytes(b"x")
    now = time.time()
    os.utime(older, (now - 3600, now - 3600))
    os.utime(newer, (now, now))
    listing = actions.search_files("invoice", ext="pdf")
    lines = [line for line in listing.splitlines() if line.startswith("- ")]
    assert [Path(line[2:].split(" (")[0]).name for line in lines] == ["invoice-april.pdf", "invoice-march.pdf"]
    assert "invoice-april.pdf" in actions.search_files("", ext=".pdf")
